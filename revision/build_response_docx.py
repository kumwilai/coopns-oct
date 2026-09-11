#!/usr/bin/env python3
"""Build the response to reviewers as a Word document.

Content lives in responses.py as a plain data structure so the wording can be
edited without touching the layout code.

Run with /home/kumwilai/osmnx-env/bin/python revision/build_response_docx.py
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from responses import META, OPENING, SUMMARY_OF_CHANGES, REVIEWERS, CLOSING  # noqa: E402

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
QUOTE_BG = "EEF3F8"
GRAY = RGBColor(0x44, 0x44, 0x44)


def shade(paragraph, hexcolor):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    ppr.append(shd)


def border(paragraph, side="left", size=18, color="1F4E79"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "8")
    el.set(qn("w:color"), color)
    borders.append(el)
    ppr.append(borders)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)


def add_title(doc, text, size=16):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(10)
    return p


def add_heading(doc, text, size=13, space_before=14):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, italic=False, bold=False, indent=0.0, size=10.5, color=None):
    for block in [b for b in text.strip().split("\n\n") if b.strip()]:
        p = doc.add_paragraph()
        r = p.add_run(" ".join(block.split()))
        r.italic = italic
        r.bold = bold
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        p.paragraph_format.left_indent = Inches(indent)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc


def add_comment_block(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    shade(p, QUOTE_BG)
    border(p, "left")
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = ACCENT
    r2 = p.add_run(" ".join(text.split()))
    r2.italic = True
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = GRAY


def add_response(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("Response")
    r.bold = True
    r.font.size = Pt(10.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    add_body(doc, text)


def add_changes(doc, items):
    if not items:
        return
    p = doc.add_paragraph()
    r = p.add_run("Changes in the manuscript")
    r.bold = True
    r.font.size = Pt(10.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    for it in items:
        b = doc.add_paragraph(style="List Bullet")
        run = b.add_run(" ".join(it.split()))
        run.font.size = Pt(10.5)
        b.paragraph_format.space_after = Pt(2)


def add_table(doc, caption, header, rows):
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.bold = True
        r.font.size = Pt(9.5)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
    return t


def build():
    doc = Document()
    setup_styles(doc)

    add_title(doc, "Response to Reviewers")
    add_body(doc, META["manuscript"], bold=True)
    add_body(doc, META["title"], italic=True)
    add_body(doc, META["date"])

    add_heading(doc, "Opening")
    add_body(doc, OPENING)

    add_heading(doc, "Summary of the main changes")
    for i, item in enumerate(SUMMARY_OF_CHANGES, 1):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(" ".join(item.split()))
        run.font.size = Pt(10.5)

    for rev in REVIEWERS:
        doc.add_page_break()
        add_heading(doc, rev["name"], size=14, space_before=0)
        if rev.get("preamble"):
            add_body(doc, rev["preamble"])
        for c in rev["comments"]:
            add_comment_block(doc, c["label"], c["comment"])
            add_response(doc, c["response"])
            if c.get("table"):
                add_table(doc, c["table"].get("caption"),
                          c["table"]["header"], c["table"]["rows"])
                if c["table"].get("note"):
                    add_body(doc, c["table"]["note"], size=9)
            add_changes(doc, c.get("changes", []))

    doc.add_page_break()
    add_heading(doc, "Closing", size=14, space_before=0)
    add_body(doc, CLOSING)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "Response_to_Reviewers.docx")
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build()
